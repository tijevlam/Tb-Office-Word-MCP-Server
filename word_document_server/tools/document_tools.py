"""
Document creation and manipulation tools for Word Document Server.
"""
import os
import json
import logging
from io import BytesIO
from typing import Dict, List, Optional, Any
from zipfile import ZIP_DEFLATED, ZipFile
from docx import Document

from word_document_server.utils.file_utils import check_file_writeable, ensure_docx_extension, create_document_copy
from word_document_server.utils.document_utils import get_document_properties, extract_document_text, get_document_structure, get_document_xml, insert_header_near_text, insert_line_or_paragraph_near_text
from word_document_server.core.styles import ensure_heading_style, ensure_table_style

logger = logging.getLogger(__name__)


def _load_template(template_filename: str):
    extension = os.path.splitext(template_filename)[1].lower()
    if extension == ".docx":
        return Document(template_filename)
    if extension != ".dotx":
        raise ValueError("Template must be a .docx or .dotx file")

    template_content_type = b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
    document_content_type = b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"

    with ZipFile(template_filename, "r") as template:
        content_types = template.read("[Content_Types].xml")
        if template_content_type not in content_types:
            raise ValueError("DOTX template has an unexpected document content type")

        converted_package = BytesIO()
        with ZipFile(converted_package, "w", ZIP_DEFLATED) as converted:
            converted.comment = template.comment
            for entry in template.infolist():
                content = template.read(entry.filename)
                if entry.filename == "[Content_Types].xml":
                    content = content.replace(template_content_type, document_content_type)
                converted.writestr(entry, content)

    return Document(BytesIO(converted_package.getvalue()))


TEMPLATE_ENV_VAR = "WORD_DOCUMENT_TEMPLATE"
TEMPLATE_EXTENSIONS = (".docx", ".dotx")


def _template_env_entries() -> List[str]:
    """Entries of WORD_DOCUMENT_TEMPLATE: files and/or directories separated by os.pathsep."""
    raw = os.getenv(TEMPLATE_ENV_VAR, "")
    entries = [os.path.expanduser(p.strip()) for p in raw.split(os.pathsep) if p.strip()]
    logger.debug("%s=%r -> %d entries", TEMPLATE_ENV_VAR, raw, len(entries))
    for entry in entries:
        kind = "directory" if os.path.isdir(entry) else "file" if os.path.isfile(entry) else "NOT FOUND"
        logger.debug("  template entry %s: %s", entry, kind)
    return entries


def _template_directories() -> List[str]:
    return [p for p in _template_env_entries() if os.path.isdir(p)]


def _list_templates() -> List[str]:
    """All template files reachable through WORD_DOCUMENT_TEMPLATE."""
    found: List[str] = []
    for entry in _template_env_entries():
        if os.path.isfile(entry):
            found.append(entry)
        elif os.path.isdir(entry):
            for name in sorted(os.listdir(entry)):
                if name.startswith("~$"):
                    continue
                if name.lower().endswith(TEMPLATE_EXTENSIONS):
                    found.append(os.path.join(entry, name))
                else:
                    logger.debug("  skipping %s (not .docx/.dotx)", os.path.join(entry, name))
    logger.debug("Templates found: %s", found)
    return found


def _resolve_template(template_filename: Optional[str]) -> Optional[str]:
    """Resolve a template name/path to a file, or None when no template applies.

    Without a name, the first file entry of WORD_DOCUMENT_TEMPLATE is the default.
    With a name, an existing path wins; otherwise it is looked up (with or without
    extension) in the configured template directories and single-file entries.
    """
    if not template_filename:
        for entry in _template_env_entries():
            if os.path.isfile(entry):
                logger.debug("No template requested; using default %s", entry)
                return entry
        logger.debug("No template requested and no default file configured; using blank document")
        return None

    logger.debug("Resolving template %r", template_filename)
    if os.path.isfile(template_filename):
        logger.debug("  %r is an existing path", template_filename)
        return template_filename

    wanted = os.path.basename(template_filename).lower()
    for candidate in _list_templates():
        base = os.path.basename(candidate).lower()
        if wanted in (base, os.path.splitext(base)[0]):
            logger.debug("  matched %r to %s", template_filename, candidate)
            return candidate
    logger.debug("  no template matched %r", template_filename)
    return template_filename


async def create_document(
    filename: str,
    title: Optional[str] = None,
    author: Optional[str] = None,
    template_filename: Optional[str] = None,
) -> str:
    """Create a new Word document with optional metadata.
    
    Args:
        filename: Name of the document to create (with or without .docx extension)
        title: Optional title for the document metadata
        author: Optional author for the document metadata
        template_filename: Optional template: a path, or the name of a template found in
            the WORD_DOCUMENT_TEMPLATE directories (see list_templates)
    """
    filename = ensure_docx_extension(filename)
    template_filename = _resolve_template(template_filename)
    
    # Check if file is writeable
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        logger.error("Cannot create document %s: %s", filename, error_message)
        return f"Cannot create document: {error_message}"
    
    try:
        if template_filename:
            if not os.path.isfile(template_filename):
                logger.error("Template %s does not exist", template_filename)
                available = [os.path.basename(t) for t in _list_templates()]
                hint = f". Available templates: {', '.join(available)}" if available else ""
                return f"Template {template_filename} does not exist{hint}"
            doc = _load_template(template_filename)
        else:
            doc = Document()
        
        # Set properties if provided
        if title:
            doc.core_properties.title = title
        if author:
            doc.core_properties.author = author
        
        # Ensure necessary styles exist
        ensure_heading_style(doc)
        ensure_table_style(doc)
        
        # Save the document
        doc.save(filename)
        
        return f"Document {filename} created successfully"
    except Exception as e:
        logger.exception("Failed to create document %s (template: %s)", filename, template_filename)
        return f"Failed to create document: {str(e)}"


async def list_templates() -> str:
    """List the templates available through the WORD_DOCUMENT_TEMPLATE environment variable."""
    if not _template_env_entries():
        return f"No templates configured. Set {TEMPLATE_ENV_VAR} to a template file or a directory."
    templates = _list_templates()
    if not templates:
        return f"No .docx or .dotx templates found via {TEMPLATE_ENV_VAR}"
    default = _resolve_template(None)
    lines = [f"Found {len(templates)} templates:"]
    for t in templates:
        lines.append(f"- {os.path.basename(t)} ({t})" + (" [default]" if t == default else ""))
    return "\n".join(lines)


async def get_document_info(filename: str) -> str:
    """Get information about a Word document.
    
    Args:
        filename: Path to the Word document
    """
    filename = ensure_docx_extension(filename)
    
    if not os.path.exists(filename):
        return f"Document {filename} does not exist"
    
    try:
        properties = get_document_properties(filename)
        return json.dumps(properties, indent=2)
    except Exception as e:
        return f"Failed to get document info: {str(e)}"


async def get_document_text(filename: str) -> str:
    """Extract all text from a Word document.
    
    Args:
        filename: Path to the Word document
    """
    filename = ensure_docx_extension(filename)
    
    return extract_document_text(filename)


async def get_document_outline(filename: str) -> str:
    """Get the structure of a Word document.
    
    Args:
        filename: Path to the Word document
    """
    filename = ensure_docx_extension(filename)
    
    structure = get_document_structure(filename)
    return json.dumps(structure, indent=2)


async def list_available_documents(directory: str = ".") -> str:
    """List all .docx files in the specified directory.
    
    Args:
        directory: Directory to search for Word documents
    """
    try:
        if not os.path.exists(directory):
            return f"Directory {directory} does not exist"
        
        docx_files = [f for f in os.listdir(directory) if f.endswith('.docx')]
        
        if not docx_files:
            return f"No Word documents found in {directory}"
        
        result = f"Found {len(docx_files)} Word documents in {directory}:\n"
        for file in docx_files:
            file_path = os.path.join(directory, file)
            size = os.path.getsize(file_path) / 1024  # KB
            result += f"- {file} ({size:.2f} KB)\n"
        
        return result
    except Exception as e:
        return f"Failed to list documents: {str(e)}"


async def copy_document(source_filename: str, destination_filename: Optional[str] = None) -> str:
    """Create a copy of a Word document.
    
    Args:
        source_filename: Path to the source document
        destination_filename: Optional path for the copy. If not provided, a default name will be generated.
    """
    source_filename = ensure_docx_extension(source_filename)
    
    if destination_filename:
        destination_filename = ensure_docx_extension(destination_filename)
    
    success, message, new_path = create_document_copy(source_filename, destination_filename)
    if success:
        return message
    else:
        return f"Failed to copy document: {message}"


async def merge_documents(target_filename: str, source_filenames: List[str], add_page_breaks: bool = True) -> str:
    """Merge multiple Word documents into a single document.
    
    Args:
        target_filename: Path to the target document (will be created or overwritten)
        source_filenames: List of paths to source documents to merge
        add_page_breaks: If True, add page breaks between documents
    """
    from word_document_server.core.tables import copy_table
    
    target_filename = ensure_docx_extension(target_filename)
    
    # Check if target file is writeable
    is_writeable, error_message = check_file_writeable(target_filename)
    if not is_writeable:
        return f"Cannot create target document: {error_message}"
    
    # Validate all source documents exist
    missing_files = []
    for filename in source_filenames:
        doc_filename = ensure_docx_extension(filename)
        if not os.path.exists(doc_filename):
            missing_files.append(doc_filename)
    
    if missing_files:
        return f"Cannot merge documents. The following source files do not exist: {', '.join(missing_files)}"
    
    try:
        # Create a new document for the merged result
        target_doc = Document()
        
        # Process each source document
        for i, filename in enumerate(source_filenames):
            doc_filename = ensure_docx_extension(filename)
            source_doc = Document(doc_filename)
            
            # Add page break between documents (except before the first one)
            if add_page_breaks and i > 0:
                target_doc.add_page_break()
            
            # Copy all paragraphs
            for paragraph in source_doc.paragraphs:
                # Create a new paragraph with the same text and style
                new_paragraph = target_doc.add_paragraph(paragraph.text)
                new_paragraph.style = target_doc.styles['Normal']  # Default style
                
                # Try to match the style if possible
                try:
                    if paragraph.style and paragraph.style.name in target_doc.styles:
                        new_paragraph.style = target_doc.styles[paragraph.style.name]
                except:
                    pass
                
                # Copy run formatting
                for i, run in enumerate(paragraph.runs):
                    if i < len(new_paragraph.runs):
                        new_run = new_paragraph.runs[i]
                        # Copy basic formatting
                        new_run.bold = run.bold
                        new_run.italic = run.italic
                        new_run.underline = run.underline
                        # Font size if specified
                        if run.font.size:
                            new_run.font.size = run.font.size
            
            # Copy all tables
            for table in source_doc.tables:
                copy_table(table, target_doc)
        
        # Save the merged document
        target_doc.save(target_filename)
        return f"Successfully merged {len(source_filenames)} documents into {target_filename}"
    except Exception as e:
        return f"Failed to merge documents: {str(e)}"


async def get_document_xml_tool(filename: str) -> str:
    """Get the raw XML structure of a Word document."""
    return get_document_xml(filename)
