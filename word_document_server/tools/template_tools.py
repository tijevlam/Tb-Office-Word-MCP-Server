"""
Tools to fill Word templates in place (title page, headers/footers, placeholders).
"""
import json
import logging
import os
from typing import Dict

from docx import Document

from word_document_server.core.template_fill import fill_template as _fill_template, list_fields
from word_document_server.utils.file_utils import check_file_writeable, ensure_docx_extension

logger = logging.getLogger(__name__)


async def list_template_fields(filename: str) -> str:
    """Show what can be filled in a document created from a template.

    Lists content controls (tag/alias/current text), placeholder-like text such as
    {{title}} or [Client], text boxes, and header/footer text.

    Args:
        filename: Path to the Word document
    """
    filename = ensure_docx_extension(filename)
    if not os.path.exists(filename):
        return f"Document {filename} does not exist"
    try:
        fields = list_fields(Document(filename))
        logger.debug("Template fields in %s: %s", filename, fields)
        return json.dumps(fields, indent=2, ensure_ascii=False)
    except Exception as e:
        logger.exception("Failed to list template fields in %s", filename)
        return f"Failed to list template fields: {str(e)}"


async def fill_template(filename: str, replacements: Dict[str, str]) -> str:
    """Fill a template in place instead of appending text at the end.

    Each key is matched against content control tags/aliases (case-insensitive) and
    against literal text anywhere in the document: body, tables, headers, footers and
    text boxes, also when the text is split over several runs.

    Args:
        filename: Path to the Word document
        replacements: Mapping of placeholder/content-control name to the new text
    """
    filename = ensure_docx_extension(filename)
    if not os.path.exists(filename):
        return f"Document {filename} does not exist"
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        return f"Cannot modify document: {error_message}. Consider creating a copy first."
    try:
        doc = Document(filename)
        report = _fill_template(doc, replacements)
        doc.save(filename)
        logger.debug("fill_template %s: %s", filename, report)
        return json.dumps(report, indent=2, ensure_ascii=False)
    except Exception as e:
        logger.exception("Failed to fill template %s", filename)
        return f"Failed to fill template: {str(e)}"
