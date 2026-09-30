"""
Fill Word templates in place: content controls, placeholder text, headers/footers and text boxes.

python-docx only exposes plain body paragraphs, so title pages built from content controls
(w:sdt) or text boxes are invisible to it. These helpers work on the XML of every story
(body, headers, footers) instead.
"""
import logging
import re
from typing import Any, Dict, Iterator, List, Optional, Tuple

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

logger = logging.getLogger(__name__)

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
PLACEHOLDER_RE = re.compile(r"\{\{[^{}]+\}\}|<<[^<>]+>>|\$\{[^{}]+\}|\[[^\[\]]{1,60}\]")
CORE_PROPERTY_FOR_BINDING = {
    "title": "title",
    "creator": "author",
    "subject": "subject",
    "description": "comments",
    "keywords": "keywords",
}


def iter_stories(doc) -> Iterator[Tuple[str, Any]]:
    """Yield (location, root element) for the body and every own header/footer."""
    yield "body", doc.element.body
    seen = set()
    for number, section in enumerate(doc.sections, start=1):
        for kind in (
            "header", "first_page_header", "even_page_header",
            "footer", "first_page_footer", "even_page_footer",
        ):
            part = getattr(section, kind)
            if part.is_linked_to_previous:
                continue
            element = part._element
            if id(element) in seen:
                continue
            seen.add(id(element))
            yield f"{kind.replace('_', ' ')} (section {number})", element


def _owner_paragraph(element):
    parent = element.getparent()
    while parent is not None and parent.tag != qn("w:p"):
        parent = parent.getparent()
    return parent


def _own_texts(paragraph) -> List[Any]:
    """w:t elements that belong to this paragraph (not to a paragraph in a nested text box)."""
    return [t for t in paragraph.iter(qn("w:t")) if _owner_paragraph(t) is paragraph]


def _paragraph_text(paragraph) -> str:
    return "".join(t.text or "" for t in _own_texts(paragraph))


def _set_text(t_element, text: str) -> None:
    t_element.text = text
    t_element.set(XML_SPACE, "preserve")


def _in_text_box(paragraph) -> bool:
    parent = paragraph.getparent()
    while parent is not None:
        if parent.tag == qn("w:txbxContent"):
            return True
        parent = parent.getparent()
    return False


def _replace_in_paragraph(paragraph, replacements: Dict[str, str]) -> Dict[str, int]:
    """Replace keys in one paragraph, also when a key is split over several runs."""
    counts: Dict[str, int] = {}
    texts = _own_texts(paragraph)
    if not texts:
        return counts
    for key, value in replacements.items():
        if not key:
            continue
        position = 0
        while True:
            full = "".join(t.text or "" for t in texts)
            index = full.find(key, position)
            if index < 0:
                break
            end = index + len(key)
            offset = 0
            first = True
            for t in texts:
                text = t.text or ""
                start, stop = offset, offset + len(text)
                offset = stop
                if stop <= index or start >= end:
                    continue
                low = max(index, start) - start
                high = min(end, stop) - start
                _set_text(t, text[:low] + (value if first else "") + text[high:])
                first = False
            counts[key] = counts.get(key, 0) + 1
            position = index + len(value)
    return counts


def _sdt_identity(sdt) -> Tuple[Optional[str], Optional[str]]:
    properties = sdt.find(qn("w:sdtPr"))
    if properties is None:
        return None, None
    tag = properties.find(qn("w:tag"))
    alias = properties.find(qn("w:alias"))
    return (
        tag.get(qn("w:val")) if tag is not None else None,
        alias.get(qn("w:val")) if alias is not None else None,
    )


def _sdt_is_placeholder(sdt) -> bool:
    properties = sdt.find(qn("w:sdtPr"))
    return properties is not None and properties.find(qn("w:showingPlcHdr")) is not None


def _sdt_text(sdt) -> str:
    content = sdt.find(qn("w:sdtContent"))
    if content is None:
        return ""
    return "".join(t.text or "" for t in content.iter(qn("w:t")))


def _bound_core_property(sdt) -> Optional[str]:
    properties = sdt.find(qn("w:sdtPr"))
    binding = properties.find(qn("w:dataBinding")) if properties is not None else None
    if binding is None:
        return None
    xpath = binding.get(qn("w:xpath")) or ""
    if "coreProperties" not in xpath:
        return None
    match = re.findall(r"/(?:\w+:)?(\w+)\[", xpath)
    return CORE_PROPERTY_FOR_BINDING.get(match[-1]) if match else None


def _fill_sdt(doc, sdt, value: str) -> None:
    content = sdt.find(qn("w:sdtContent"))
    if content is None:
        return
    properties = sdt.find(qn("w:sdtPr"))

    texts = list(content.iter(qn("w:t")))
    if texts:
        _set_text(texts[0], value)
        for t in texts[1:]:
            t.text = ""
    else:
        container = content.find(qn("w:p"))
        if container is None:
            container = content
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        _set_text(t, value)
        run.append(t)
        container.append(run)

    # Drop the grey placeholder look now that real text is present.
    if properties is not None:
        showing = properties.find(qn("w:showingPlcHdr"))
        if showing is not None:
            properties.remove(showing)
    for run_style in list(content.iter(qn("w:rStyle"))):
        if run_style.get(qn("w:val")) == "PlaceholderText":
            run_style.getparent().remove(run_style)

    core_property = _bound_core_property(sdt)
    if core_property:
        setattr(doc.core_properties, core_property, value)
    elif properties is not None:
        # Other bindings would make Word overwrite the typed text with the stored value.
        binding = properties.find(qn("w:dataBinding"))
        if binding is not None:
            properties.remove(binding)


def fill_template(doc, replacements: Dict[str, str]) -> Dict[str, Any]:
    """Fill content controls (by tag/alias) and replace text everywhere. Returns a report."""
    wanted = {key.lower(): key for key in replacements if key}
    controls: Dict[str, int] = {}
    text_counts: Dict[str, int] = {}

    for _, root in iter_stories(doc):
        for sdt in root.iter(qn("w:sdt")):
            for identity in _sdt_identity(sdt):
                key = wanted.get((identity or "").lower())
                if key:
                    _fill_sdt(doc, sdt, replacements[key])
                    controls[key] = controls.get(key, 0) + 1
                    break

    for _, root in iter_stories(doc):
        for paragraph in root.iter(qn("w:p")):
            for key, count in _replace_in_paragraph(paragraph, replacements).items():
                text_counts[key] = text_counts.get(key, 0) + count

    not_found = [k for k in replacements if k not in controls and k not in text_counts]
    return {"content_controls_filled": controls, "text_replaced": text_counts, "not_found": not_found}


def list_fields(doc) -> Dict[str, Any]:
    """Describe what can be filled: content controls, placeholder-like text, text boxes, headers/footers."""
    controls: List[Dict[str, Any]] = []
    placeholders: List[Dict[str, str]] = []
    text_boxes: List[Dict[str, str]] = []
    headers_footers: List[Dict[str, str]] = []
    seen = set()

    for location, root in iter_stories(doc):
        for sdt in root.iter(qn("w:sdt")):
            tag, alias = _sdt_identity(sdt)
            controls.append({
                "location": location,
                "tag": tag,
                "alias": alias,
                "current_text": _sdt_text(sdt),
                "is_empty_placeholder": _sdt_is_placeholder(sdt),
                "bound_to_document_property": _bound_core_property(sdt),
            })
        for paragraph in root.iter(qn("w:p")):
            text = _paragraph_text(paragraph).strip()
            if not text:
                continue
            if location != "body" and ("hf", location, text) not in seen:
                seen.add(("hf", location, text))
                headers_footers.append({"location": location, "text": text})
            if _in_text_box(paragraph) and ("tb", location, text) not in seen:
                seen.add(("tb", location, text))
                text_boxes.append({"location": location, "text": text})
            for match in PLACEHOLDER_RE.findall(text):
                if ("ph", location, match) not in seen:
                    seen.add(("ph", location, match))
                    placeholders.append({"location": location, "text": match})

    return {
        "content_controls": controls,
        "placeholder_like_text": placeholders,
        "text_boxes": text_boxes,
        "headers_footers": headers_footers,
    }
