import asyncio
import json
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from word_document_server.tools.template_tools import fill_template, list_template_fields


def _sdt(tag: str, text: str, placeholder: bool = False) -> str:
    showing = "<w:showingPlcHdr/>" if placeholder else ""
    style = '<w:rPr><w:rStyle w:val="PlaceholderText"/></w:rPr>' if placeholder else ""
    return (
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:alias w:val="{tag} alias"/><w:tag w:val="{tag}"/>'
        f"{showing}</w:sdtPr><w:sdtContent><w:p><w:r>{style}<w:t>{text}</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
    )


def _make_template(path: Path) -> None:
    doc = Document()
    body = doc.element.body
    body.insert(0, parse_xml(_sdt("title", "Klik om titel in te voeren", placeholder=True)))

    split = doc.add_paragraph()
    split.add_run("Client: {{cl")
    split.add_run("ient}} done")

    doc.add_paragraph("Closing paragraph")
    doc.sections[0].header.paragraphs[0].text = "Header {{client}}"

    textbox = parse_xml(
        f'<w:p {nsdecls("w")}><w:r><w:pict><w:txbxContent><w:p><w:r><w:t>[Subtitel]</w:t></w:r></w:p>'
        "</w:txbxContent></w:pict></w:r></w:p>"
    )
    body.insert(1, textbox)
    doc.save(path)


def test_list_template_fields_finds_hidden_content(tmp_path: Path):
    path = tmp_path / "t.docx"
    _make_template(path)

    fields = json.loads(asyncio.run(list_template_fields(str(path))))

    assert fields["content_controls"][0]["tag"] == "title"
    assert fields["content_controls"][0]["is_empty_placeholder"] is True
    assert {"location": "body", "text": "{{client}}"} in fields["placeholder_like_text"]
    assert any(t["text"] == "[Subtitel]" for t in fields["text_boxes"])
    assert any("Header" in h["text"] for h in fields["headers_footers"])


def test_fill_template_replaces_in_place(tmp_path: Path):
    path = tmp_path / "t.docx"
    _make_template(path)

    report = json.loads(
        asyncio.run(
            fill_template(
                str(path),
                {"title": "Jaarrapport", "{{client}}": "Acme", "[Subtitel]": "Q3", "missing": "x"},
            )
        )
    )

    doc = Document(path)
    xml = doc.element.xml
    assert report["content_controls_filled"] == {"title": 1}
    assert report["text_replaced"]["{{client}}"] == 2
    assert report["not_found"] == ["missing"]
    assert "Jaarrapport" in xml and "Klik om titel" not in xml
    assert "PlaceholderText" not in xml and "showingPlcHdr" not in xml
    assert "Q3" in xml
    assert doc.paragraphs[0].text == "Client: Acme done" or "Client: Acme done" in [p.text for p in doc.paragraphs]
    assert doc.paragraphs[-1].text == "Closing paragraph"
    assert doc.sections[0].header.paragraphs[0].text == "Header Acme"
