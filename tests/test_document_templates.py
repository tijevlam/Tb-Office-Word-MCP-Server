import asyncio
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from word_document_server.tools.document_tools import create_document


def _make_template(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Corporate template content")
    doc.styles["Normal"].font.name = "Aptos"
    doc.save(path)


def _make_dotx_template(docx_path: Path, dotx_path: Path) -> None:
    document_content_type = b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
    template_content_type = b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"

    with ZipFile(docx_path, "r") as source, ZipFile(dotx_path, "w", ZIP_DEFLATED) as template:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "[Content_Types].xml":
                content = content.replace(document_content_type, template_content_type)
            template.writestr(entry, content)


@pytest.mark.parametrize("template_extension", [".docx", ".dotx"])
def test_create_document_from_template_preserves_content_and_styles(
    tmp_path: Path, template_extension: str
):
    template_docx = tmp_path / "corporate-template.docx"
    _make_template(template_docx)
    template_path = template_docx
    if template_extension == ".dotx":
        template_path = tmp_path / "corporate-template.dotx"
        _make_dotx_template(template_docx, template_path)

    output_path = tmp_path / "created-document"
    result = asyncio.run(
        create_document(
            str(output_path),
            title="Template-based document",
            author="Test author",
            template_filename=str(template_path),
        )
    )

    created_document = Document(output_path.with_suffix(".docx"))
    assert result == f"Document {output_path}.docx created successfully"
    assert created_document.paragraphs[0].text == "Corporate template content"
    assert created_document.styles["Normal"].font.name == "Aptos"
    assert created_document.core_properties.title == "Template-based document"
    assert created_document.core_properties.author == "Test author"


def test_create_document_without_template_remains_blank(tmp_path: Path, monkeypatch):
    monkeypatch.delenv("WORD_DOCUMENT_TEMPLATE", raising=False)
    output_path = tmp_path / "blank.docx"

    result = asyncio.run(create_document(str(output_path)))

    assert result == f"Document {output_path} created successfully"
    assert Document(output_path).paragraphs == []


def test_create_document_uses_environment_template_by_default(
    tmp_path: Path, monkeypatch
):
    template_path = tmp_path / "default-template.docx"
    _make_template(template_path)
    monkeypatch.setenv("WORD_DOCUMENT_TEMPLATE", str(template_path))
    output_path = tmp_path / "created-from-default.docx"

    result = asyncio.run(create_document(str(output_path)))

    assert result == f"Document {output_path} created successfully"
    assert Document(output_path).paragraphs[0].text == "Corporate template content"


def test_explicit_template_overrides_environment_template(tmp_path: Path, monkeypatch):
    default_template = tmp_path / "default-template.docx"
    explicit_template = tmp_path / "explicit-template.docx"
    _make_template(default_template)
    explicit_doc = Document()
    explicit_doc.add_paragraph("Explicit template content")
    explicit_doc.save(explicit_template)
    monkeypatch.setenv("WORD_DOCUMENT_TEMPLATE", str(default_template))
    output_path = tmp_path / "created-from-explicit.docx"

    result = asyncio.run(
        create_document(str(output_path), template_filename=str(explicit_template))
    )

    assert result == f"Document {output_path} created successfully"
    assert Document(output_path).paragraphs[0].text == "Explicit template content"


def test_create_document_rejects_missing_template(tmp_path: Path):
    output_path = tmp_path / "created.docx"
    template_path = tmp_path / "missing.docx"

    result = asyncio.run(
        create_document(str(output_path), template_filename=str(template_path))
    )

    assert result == f"Template {template_path} does not exist"
    assert not output_path.exists()


def test_create_document_rejects_unsupported_template_extension(tmp_path: Path):
    template_path = tmp_path / "template.dot"
    template_path.touch()

    result = asyncio.run(
        create_document(
            str(tmp_path / "created.docx"), template_filename=str(template_path)
        )
    )

    assert "Template must be a .docx or .dotx file" in result


def _make_named_template(path: Path, text: str) -> None:
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)


def test_template_directory_resolves_template_by_name(tmp_path: Path, monkeypatch):
    template_dir = tmp_path / "templates"
    template_dir.mkdir()
    _make_named_template(template_dir / "report.docx", "Report template")
    _make_named_template(template_dir / "letter.docx", "Letter template")
    monkeypatch.setenv("WORD_DOCUMENT_TEMPLATE", str(template_dir))
    output_path = tmp_path / "out.docx"

    result = asyncio.run(create_document(str(output_path), template_filename="letter"))

    assert result == f"Document {output_path} created successfully"
    assert Document(output_path).paragraphs[0].text == "Letter template"


def test_template_directory_without_name_stays_blank(tmp_path: Path, monkeypatch):
    template_dir = tmp_path / "templates"
    template_dir.mkdir()
    _make_named_template(template_dir / "report.docx", "Report template")
    monkeypatch.setenv("WORD_DOCUMENT_TEMPLATE", str(template_dir))
    output_path = tmp_path / "out.docx"

    asyncio.run(create_document(str(output_path)))

    assert Document(output_path).paragraphs == []


def test_list_templates_shows_directory_contents(tmp_path: Path, monkeypatch):
    from word_document_server.tools.document_tools import list_templates

    template_dir = tmp_path / "templates"
    template_dir.mkdir()
    _make_named_template(template_dir / "report.docx", "Report template")
    monkeypatch.setenv("WORD_DOCUMENT_TEMPLATE", str(template_dir))

    result = asyncio.run(list_templates())

    assert "report.docx" in result


def test_debug_logging_shows_search_and_found_templates(tmp_path: Path, monkeypatch, caplog):
    import logging

    template_dir = tmp_path / "templates"
    template_dir.mkdir()
    _make_named_template(template_dir / "report.docx", "Report template")
    monkeypatch.setenv("WORD_DOCUMENT_TEMPLATE", str(template_dir))

    with caplog.at_level(logging.DEBUG):
        asyncio.run(create_document(str(tmp_path / "out.docx"), template_filename="report"))

    assert "directory" in caplog.text
    assert "matched 'report'" in caplog.text
